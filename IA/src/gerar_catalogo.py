"""Etapa 4: catálogo padronizado (XLSX formatado), um item por código de serviço canônico."""
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import config
from dados import carregar_base_saneada, ids_servico
from normalizacao import limpar_espacos

CABECALHOS = [
    ("ID Canônico (CS)", 16), ("Grupo de Serviço", 18), ("Descrição Padronizada", 45),
    ("Descrição Normalizada", 35), ("Unidade de Medida", 18), ("Qtd Caracteres", 16),
    ("Total Ocorrências", 18), ("Códigos Legados Associados", 40),
]
COLUNAS_CENTRALIZADAS = {1, 5, 6, 7}
MAX_LEGADOS_EXIBIDOS = 5


def montar_catalogo(df):
    map_cs = ids_servico(df)
    df = df.assign(csId=df["DescricaoNormalizada"].map(map_cs))
    itens = []
    for cs_id, grupo in df.groupby("csId", sort=True):
        r = grupo.iloc[0]
        descricao = limpar_espacos(r["DescricaoOriginal"])
        legados = list(dict.fromkeys(grupo["Código do Serviço"].astype(str)))
        texto_legados = ", ".join(legados[:MAX_LEGADOS_EXIBIDOS]) + ("..." if len(legados) > MAX_LEGADOS_EXIBIDOS else "")
        itens.append([cs_id, r["Grupo de Serviço"], descricao, r["DescricaoNormalizada"],
                      r["Unidade medida"], len(descricao), len(grupo), texto_legados])
    return itens


def salvar_planilha(itens):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catálogo Padronizado"
    borda = Border(*(Side(style="thin", color="D5DBDB"),) * 4)
    zebra = PatternFill(start_color="F8F9F9", end_color="F8F9F9", fill_type="solid")

    ws.merge_cells("A1:H1")
    ws["A1"] = "GERDAU — Catálogo Padronizado de Serviços (Saneado)"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="2C3E50")
    ws.merge_cells("A2:H2")
    ws["A2"] = f"Total de itens canônicos únicos: {len(itens)} | IDs iguais ao _id de codigos_servico no MongoDB"
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="7F8C8D")

    for col, (titulo, largura) in enumerate(CABECALHOS, start=1):
        cell = ws.cell(row=4, column=col, value=titulo)
        cell.fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = borda
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = largura

    for linha, item in enumerate(itens, start=5):
        for col, valor in enumerate(item, start=1):
            cell = ws.cell(row=linha, column=col, value=valor)
            cell.font = Font(name="Arial", size=10)
            cell.border = borda
            cell.alignment = Alignment(horizontal="center" if col in COLUNAS_CENTRALIZADAS else "left",
                                       vertical="center")
            if linha % 2 == 0:
                cell.fill = zebra

    ws.freeze_panes = "A5"
    config.PROCESSED.mkdir(parents=True, exist_ok=True)
    wb.save(config.CATALOGO)


def main():
    itens = montar_catalogo(carregar_base_saneada())
    salvar_planilha(itens)
    print(f"[OK] Catálogo com {len(itens)} itens -> {config.CATALOGO}")


if __name__ == "__main__":
    main()
