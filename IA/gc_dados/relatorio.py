"""Relatorio Excel: Resumo (formulas), Catalogo, De-Para, Pendencias, Validacao migracao, Lote, Qualidade,
Amostra p rotular, Amostra de-para, Dicionario a confirmar."""
import random
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .normalizador import carregar_dicionario

F = "Arial"
CAB = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFF2CC")


def _tabela(ws, cabecalho, linhas, larguras, linha_cab=1):
    ws.append(cabecalho)
    for c in ws[ws.max_row]:
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = CAB
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for l in linhas: ws.append(l)
    for r in ws.iter_rows(min_row=linha_cab + 1):
        for c in r: c.font = Font(name=F, size=10)
    for i, w in enumerate(larguras, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(linha_cab + 1, 1).coordinate
    ws.auto_filter.ref = f"A{linha_cab}:{get_column_letter(len(cabecalho))}{ws.max_row}"


def _amostra(r, seed=7):
    rnd = random.Random(seed)
    cat = {c["_id"]: c["descricao"] for c in r["catalogo"]}
    pares = []
    vistos = set()
    for origem, dec, k in (("Fundido automaticamente (provavel)", "CONSOLIDADO_PROVAVEL", 40),
                           ("Fundido automaticamente (exato)", "CONSOLIDADO_EXATO", 30)):
        lista = [d for d in r["depara"] if d["decision"] == dec]; rnd.shuffle(lista); n = 0
        for d in lista:
            if d["legacyDescricao"] in vistos: continue
            vistos.add(d["legacyDescricao"]); n += 1
            pares.append((origem, d["legacyDescricao"], cat[d["newId"]], d["score"], d["reason"]))
            if n >= k: break
    pend = [p for p in r["pendencias"] if p["tipo"] != "CONTROLE_NEGATIVO_NUMEROS"]
    for lo, hi in ((0.75, 0.80), (0.80, 0.85), (0.85, 0.90)):
        faixa = [p for p in pend if lo <= p["score"] < hi]; rnd.shuffle(faixa)
        for p in faixa[:25]:
            pares.append((f"Alerta {lo:.2f}-{hi:.2f}", p["descricaoA"], p["descricaoB"], p["score"], p["motivo"]))
    baixa = list(r["amostra_baixa"]); rnd.shuffle(baixa)
    for p in baixa[:40]:                                   # abaixo do alerta: necessario para estimar revocacao
        pares.append(("Abaixo do alerta (0,60-0,74)", p["descricaoA"], p["descricaoB"], p["score"], p["motivo"]))
    neg = [p for p in r["pendencias"] if p["tipo"] == "CONTROLE_NEGATIVO_NUMEROS"]; rnd.shuffle(neg)
    for p in neg[:30]:
        pares.append(("Controle negativo (numeros diferentes)", p["descricaoA"], p["descricaoB"], p["score"], p["motivo"]))
    return pares


def gerar(r, caminho):
    wb = Workbook(); dic = carregar_dicionario()
    cat, lote = r["catalogo"], r["lote"]
    ws = wb.active; ws.title = "Resumo"

    wc = wb.create_sheet("Catalogo")
    _tabela(wc, ["csId", "Descricao canonica", "Descricao sugerida (<=40)", "Caracteres", "Especificidade", "Confianca",
                 "Funcao", "Modificadores", "Numeros", "Variantes", "Linhas legado", "GS", "Em revisao",
                 "Candidatos similares", "Status", "Canonico incerto", "Incompleto", "Pendencias"],
            [[c["_id"], c["descricao"], c["descricaoSugerida"], c["quantidadeCaracteres"], c["especificidade"],
              c["confiancaEspecificidade"], c["funcao"], " ".join(c["modificadores"]), " ".join(c["numeros"]),
              c["qtdVariantes"], c["qtdLinhasLegado"], " ".join(c["gsIds"]), c["emRevisao"],
              " ".join(c["candidatosSimilares"]), c["status"], c["canonicoIncerto"], c["incompleto"],
              " ".join(c["pendencias"])] for c in cat],
            [10, 46, 40, 11, 38, 11, 34, 16, 12, 10, 12, 10, 11, 24, 9, 11, 11, 26])
    wd = wb.create_sheet("De-Para")
    _tabela(wd, ["Chave legada (cs|fornecedor|um)", "Codigo servico legado", "Fornecedor legado", "Grupo", "Descricao legada",
                 "UM legada", "csId novo", "Decisao", "Motivo", "Score", "Em revisao"],
            [[d["legacyKey"], d["legacyId"], d["legacyFornecedor"], d["legacyGS"], d["legacyDescricao"], d["legacyUM"],
              d["newId"], d["decision"], d["reason"], d["score"], d["emRevisao"]] for d in r["depara"]],
            [26, 14, 14, 16, 46, 9, 10, 22, 60, 8, 11])
    wp = wb.create_sheet("Pendencias")
    _tabela(wp, ["Tipo", "csId A / item", "Descricao A / canonica", "csId B", "Descricao B", "Score", "Motivo",
                 "Decisao sugerida", "Decisao do revisor (preencher)", "Justificativa (preencher)"],
            [[c["tipo"], c.get("csIdA") or c.get("csId"), c.get("descricaoA") or c.get("descricaoCanonica"),
              c.get("csIdB"), c.get("descricaoB"), c.get("score"), c["motivo"], c.get("decisaoSugerida"), None, None]
             for c in r["controle"]],
            [34, 12, 42, 9, 42, 8, 70, 16, 26, 30])
    for row in wp.iter_rows(min_row=2, min_col=9, max_col=10):
        for c in row: c.fill = INPUT

    wv = wb.create_sheet("Validacao migracao")
    _tabela(wv, ["Secao do documento", "Verificacao", "Bloqueante", "Resultado", "Detalhe"],
            [[c["secao"], c["verificacao"], "SIM" if c["bloqueante"] else "NAO", "OK" if c["ok"] else "FALHA", c["detalhe"]]
             for c in r["checks"]], [18, 78, 12, 12, 60])

    wl = wb.create_sheet("Lote")
    pares_lote = [("Lote de importacao", lote["_id"]), ("Status", lote["status"]), ("Arquivo de origem", lote["arquivoOrigem"]),
                  ("Copia imutavel (somente leitura)", lote["copiaImutavel"]), ("SHA-256 do arquivo", lote["sha256"]),
                  ("Tamanho (bytes)", lote["tamanhoBytes"]), ("Data da importacao", lote["dataImportacao"]),
                  ("Versao do algoritmo de normalizacao", lote["versaoAlgoritmoNormalizacao"]),
                  ("Versao do dicionario", lote["versaoDicionario"]), ("Usuario/processo responsavel", lote["usuarioProcesso"]),
                  ("Resultado da validacao", "LIBERADO" if lote["resultadoValidacao"]["liberado"] else "BLOQUEADO"),
                  ("Limiar provavel / alerta", f"{lote['metricas']['limiares']['provavel']} / {lote['metricas']['limiares']['alerta']}"),
                  ("Linhas rejeitadas na entrada", lote["metricas"]["linhasRejeitadas"]),
                  ("Cobertura da matriz de-para", lote["metricas"]["coberturaDePara"])]
    _tabela(wl, ["Campo (secao 9.2)", "Valor"], [list(p) for p in pares_lote], [42, 90])
    linha_status = 3                                         # 'Status' na linha 3 da aba Lote

    wq = wb.create_sheet("Qualidade")
    _tabela(wq, ["Tipo", "Verificacao", "Ocorrencias", "Observacao"], [list(q) for q in r["qualidade"]], [14, 70, 13, 90])

    # ---- Resumo (formulas) ----
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Padronizacao dos Codigos de Servico - Resultado do saneamento (IA e Dados)"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = "Todos os numeros abaixo sao formulas sobre as demais abas."
    ws["A2"].font = Font(name=F, italic=True, size=9)
    L = 4
    def linha(rot, form, fmt="0", negrito=True):
        nonlocal L
        ws.cell(L, 1, rot).font = Font(name=F, size=10)
        c = ws.cell(L, 2, form); c.font = Font(name=F, size=10, bold=negrito); c.number_format = fmt
        L += 1
        return L - 1
    def titulo(t):
        nonlocal L
        L += 1; ws.cell(L, 1, t).font = Font(name=F, bold=True); L += 1
    titulo("Lote")
    linha("Status do lote", f"=Lote!B{linha_status}", "@")
    linha("Verificacoes de migracao OK", "=COUNTIF('Validacao migracao'!D:D,\"OK\")")
    linha("Verificacoes de migracao com FALHA", "=COUNTIF('Validacao migracao'!D:D,\"FALHA\")")
    titulo("Volume e reducao")
    r_lin = linha("Linhas na base legada", "=COUNTA('De-Para'!A:A)-1")
    r_desc = linha("Descricoes distintas no legado (apos limpar espacos)", "=SUM(Catalogo!J:J)")
    r_it = linha("Itens no catalogo padronizado", "=COUNTA(Catalogo!A:A)-1")
    linha("Reducao do catalogo (descricoes -> itens)", f"=1-B{r_it}/B{r_desc}", "0.0%")
    titulo("Decisoes do de-para")
    a = linha("Linhas MANTIDO", "=COUNTIF('De-Para'!H:H,\"MANTIDO\")")
    b = linha("Linhas CONSOLIDADO_EXATO", "=COUNTIF('De-Para'!H:H,\"CONSOLIDADO_EXATO\")")
    c_ = linha("Linhas CONSOLIDADO_PROVAVEL", "=COUNTIF('De-Para'!H:H,\"CONSOLIDADO_PROVAVEL\")")
    linha("Conferencia (decisoes = linhas legado)", f"=IF(SUM(B{a}:B{c_})=B{r_lin},\"OK\",\"DIVERGENTE\")", "@")
    titulo("Controle (revisao humana, secao 9.3 / 4.1)")
    for tipo in ("ALERTA", "PROVAVEL_NAO_FUNDIDO_BAIXA_COESAO", "CANONICO_INCERTO", "INCOMPLETO"):
        linha(f"Pendencias {tipo}", f"=COUNTIF(Pendencias!A:A,\"{tipo}\")")
    linha("Itens com similar pendente", "=COUNTIF(Catalogo!M:M,TRUE)")
    linha("Itens com descricao sugerida (<=40)", "=COUNTA(Catalogo!C:C)-1")
    titulo("Itens por especificidade (taxonomia da aba 'Mao de Obra')")
    for j, h in enumerate(["Especificidade", "Itens", "Conf. ALTA", "Conf. MEDIA", "Conf. BAIXA"], 1):
        c = ws.cell(L, j, h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = CAB
    L += 1; ini = L
    for e, _ in Counter(c["especificidade"] for c in cat).most_common():
        ws.cell(L, 1, e).font = Font(name=F, size=10)
        ws.cell(L, 2, f"=COUNTIF(Catalogo!E:E,A{L})")
        for j, conf in ((3, "ALTA"), (4, "MEDIA"), (5, "BAIXA")):
            ws.cell(L, j, f"=COUNTIFS(Catalogo!E:E,$A{L},Catalogo!F:F,\"{conf}\")")
        for j in range(2, 6): ws.cell(L, j).font = Font(name=F, size=10)
        L += 1
    ws.cell(L, 1, "Total").font = Font(name=F, bold=True)
    for j in range(2, 6):
        col = get_column_letter(j); ws.cell(L, j, f"=SUM({col}{ini}:{col}{L - 1})").font = Font(name=F, bold=True)
    L += 2
    ws.cell(L, 1, "Premissas e limites").font = Font(name=F, bold=True)
    for k, t in enumerate([
        "Fusao automatica so ocorre para: mesma descricao apos normalizacao (exato) ou diferenca apenas de digitacao/truncamento do SAP aos 40 caracteres (provavel).",
        "Qualquer termo adicional (NR35, PTA, ADM, numero, 'sem') impede a fusao automatica e gera pendencia para revisao humana.",
        "Limiares 0,90 / 0,75 vem do documento e NAO foram calibrados: preencher a aba 'Amostra p rotular' e rodar gc_dados.avaliacao.",
        "Catalogo de CS global (independente do Grupo de Servico), conforme o modelo de 4 componentes do documento.",
        "Stemming e usado so na busca (recall); nao e usado para fundir itens."], 1):
        ws.cell(L + k, 1, t).font = Font(name=F, size=9)
    ws.column_dimensions["A"].width = 62
    for col in "BCDE": ws.column_dimensions[col].width = 14

    # ---- amostras ----
    wa = wb.create_sheet("Amostra p rotular")
    wa["A1"] = ("Legenda: preencha a coluna G com DUPLICATA, DIFERENTE ou DUVIDA e rode: python -m gc_dados.avaliacao <este arquivo>. "
                "Exemplo: 'PINTOR - HORA NORMAL' x 'PINTOR (HORA NORMAL)' -> DUPLICATA.")
    wa["A1"].font = Font(name=F, italic=True, size=9); wa["A1"].alignment = Alignment(wrap_text=True)
    wa.merge_cells("A1:G1"); wa.row_dimensions[1].height = 32
    wa.append(["ID", "Origem do par", "Descricao A", "Descricao B", "Score", "Motivo", "Rotulo (preencher)"])
    for c in wa[2]: c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = CAB
    for i, (o, x, y, s, m) in enumerate(_amostra(r), 1): wa.append([i, o, x, y, s, m, None])
    for row in wa.iter_rows(min_row=3):
        for c in row: c.font = Font(name=F, size=10)
        row[6].fill = INPUT
    for i, w in enumerate([6, 36, 44, 44, 8, 70, 18], 1): wa.column_dimensions[get_column_letter(i)].width = w
    wa.freeze_panes = "A3"

    from .pipeline import amostra_depara_manual
    wm = wb.create_sheet("Amostra de-para")
    _tabela(wm, ["Decisao", "Codigo legado", "Grupo", "Descricao legada", "csId novo", "Descricao nova", "Score", "Motivo",
                 "Conferido (OK / ERRO / preencher)"],
            [[x["decision"], x["legacyId"], x["legacyGS"], x["legacyDescricao"], x["newId"], x["descricaoNova"], x["score"], x["reason"], None]
             for x in amostra_depara_manual(r)], [22, 14, 16, 46, 10, 46, 8, 60, 28])
    for row in wm.iter_rows(min_row=2, min_col=9, max_col=9):
        for c in row: c.fill = INPUT

    wdc = wb.create_sheet("Dicionario a confirmar")
    wdc["A1"] = "Suposicoes assumidas no dicionario (confirmar com quem conhece o negocio)"; wdc["A1"].font = Font(name=F, bold=True)
    for i, t in enumerate(dic["a_confirmar"], 2): wdc.cell(i, 1, t).font = Font(name=F, size=10)
    base = len(dic["a_confirmar"]) + 4
    wdc.cell(base, 1, "Termos curtos frequentes ainda nao tratados (candidatos a abreviacao):").font = Font(name=F, bold=True)
    conhecidos = set(dic["tokens"]) | set(dic["stopwords"])
    cont = Counter(t for c in cat for t in c["funcao"].split() if t.isalpha() and len(t) <= 4 and t not in conhecidos)
    for j, h in enumerate(["Termo", "Ocorrencias (itens)", "Significado (preencher)"], 1):
        c = wdc.cell(base + 1, j, h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = CAB
    for i, (t, n) in enumerate(cont.most_common(40), base + 2):
        wdc.cell(i, 1, t); wdc.cell(i, 2, n); wdc.cell(i, 3).fill = INPUT
        for j in (1, 2): wdc.cell(i, j).font = Font(name=F, size=10)
    wdc.column_dimensions["A"].width = 110; wdc.column_dimensions["B"].width = 20; wdc.column_dimensions["C"].width = 40
    wb.save(caminho)
