"""Medicao de precisao, revocacao e falsos bloqueios (secoes 16.3 e 18-fase 3 'calibrar score e pesos').

Entrada: aba 'Amostra p rotular' preenchida por especialistas (DUPLICATA / DIFERENTE / DUVIDA).
Positivo = DUPLICATA. Um par e 'previsto positivo' se score >= limiar.
Taxa de falsos bloqueios = DIFERENTES com score >= limiar / total de DIFERENTES.
ATENCAO: a amostra e estratificada (nao aleatoria); as metricas orientam a calibracao, nao sao taxas populacionais.
"""
import sys

from openpyxl import load_workbook

ROTULOS = {"DUPLICATA", "DIFERENTE", "DUVIDA"}


def carregar_rotulos(xlsx, aba="Amostra p rotular") -> list:
    ws = load_workbook(xlsx, data_only=True)[aba]
    pares = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] is None: continue
        rot = (str(row[6]).strip().upper() if row[6] else "").replace("DÚVIDA", "DUVIDA")
        if rot in ROTULOS:
            pares.append({"id": row[0], "origem": row[1], "a": row[2], "b": row[3], "score": float(row[4]), "rotulo": rot})
    return pares


def metricas(pares: list, limiar: float) -> dict:
    p = [x for x in pares if x["rotulo"] != "DUVIDA"]
    tp = sum(1 for x in p if x["rotulo"] == "DUPLICATA" and x["score"] >= limiar)
    fp = sum(1 for x in p if x["rotulo"] == "DIFERENTE" and x["score"] >= limiar)
    fn = sum(1 for x in p if x["rotulo"] == "DUPLICATA" and x["score"] < limiar)
    tn = sum(1 for x in p if x["rotulo"] == "DIFERENTE" and x["score"] < limiar)
    prec = tp / (tp + fp) if tp + fp else None
    rec = tp / (tp + fn) if tp + fn else None
    f1 = 2 * prec * rec / (prec + rec) if prec and rec else None
    dif = fp + tn
    return {"limiar": limiar, "tp": tp, "fp": fp, "fn": fn, "tn": tn, "precisao": prec, "revocacao": rec,
            "f1": f1, "taxaFalsosBloqueios": (fp / dif) if dif else None}


def varrer(pares: list, passo: float = 0.01, de: float = 0.50, ate: float = 1.0) -> list:
    out, l = [], de
    while l <= ate + 1e-9:
        out.append(metricas(pares, round(l, 4))); l += passo
    return out


def recomendar(pares: list, precisao_minima: float = 0.98, falsos_bloqueios_max: float = 0.05) -> dict:
    """limiar_provavel: menor limiar com precisao >= minima (fusao automatica segura).
       limiar_alerta : menor limiar cuja taxa de falsos bloqueios fica <= maxima, maximizando revocacao."""
    grade = varrer(pares)
    prov = [m for m in grade if m["precisao"] is not None and m["precisao"] >= precisao_minima]
    alerta = [m for m in grade if m["taxaFalsosBloqueios"] is not None and m["taxaFalsosBloqueios"] <= falsos_bloqueios_max]
    return {"limiarProvavel": prov[0]["limiar"] if prov else None, "limiarAlerta": alerta[0]["limiar"] if alerta else None,
            "pares": len([x for x in pares if x["rotulo"] != "DUVIDA"]),
            "aviso": "Amostra pequena/estratificada: use como orientacao, confirmar com mais pares rotulados."}


if __name__ == "__main__":
    pares = carregar_rotulos(sys.argv[1])
    if not pares:
        sys.exit("Nenhum par rotulado: preencha a coluna 'Rotulo' da aba 'Amostra p rotular'.")
    print("Limiar atual 0,90:", metricas(pares, 0.90)); print("Limiar atual 0,75:", metricas(pares, 0.75))
    print("Recomendacao:", recomendar(pares))
